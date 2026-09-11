"""
Build an Oracle-enabled Mass Upload .xlsm from WCRO Store Driven data.

LOCKED FORMAT (must stay consistent — matches what Bernie/Oracle bulk load expects):
  Shell:     .xlsm with Upload To Oracle button + Button2_Click ODBC macros
  Sheet:     Template (VBA code name Sheet1)
  B1:        Customer Name value
  B2:        Req. Delivery Date — blank unless --req-delivery-date / Teams bot
  B3:        Shipping Instr.
  B4:        Order Type — blank unless --order-type / Teams bot
  B5:        # of materials (item column count)
  B6:        Address Category (optional)
  Row 7:     ALWAYS blank (no WCRO stamp, no BEGINNING)
  Row 8:     SUM formulas H.. last item col
  Row 10:    G="Picking Notes"; item cols blank (no Substituting notes)
  Row 11:    G="Sku # "; item cols = retailer SKUs (numbers/text, never #N/A)
  Row 12:    G="Item"; item cols = plain descriptions (never VLOOKUP / #N/A)
             VBA skips row 12 on upload, but #N/A here still breaks usability
  Row 13:    A-D = Rep / Store # / City / PO#; E+F ALWAYS blank; G="Internal # ";
             H.. = Everde item codes
             Store # ALWAYS 4-digit zero-padded text (625 → 0625); PO# NEEDSPO/{store}
  Row 14+:   store grid; E+F ALWAYS blank; qty only where non-zero

Filters (MASS-UPLOAD ONLY — pre-upload safety):
  - Exclude citrus (item code CIT* / desc CITRUS) — ag restriction; reps order separately
  - Require active customer xref SKU×Item for the region address categories
  - Drop duplicate retailer SKUs; drop blank/#N/A descriptions; 4-digit stores

  SCOPE: These filters apply ONLY when building the Oracle .xlsm via this script /
  Teams "generate a mass upload". They must NOT be applied to WCRO extract
  (extract_wcro.py), Blob wcro_data.json, get_wcro_dashboard, or portal WCRO UI —
  citrus and all Store Driven lines stay visible in WCRO for planning/Q&A.

Usage:
    python build_mass_upload.py --channel HD --region N.CA --order-type "WIN NORTH CA" --req-delivery-date 2026-09-30
    python build_mass_upload.py --channel LOW --region N.CA --include-citrus   # override citrus drop
    python build_mass_upload.py --skip-xref-filter                             # override xref gate
"""
from __future__ import annotations

import argparse
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


DEFAULT_ORACLE_TEMPLATE = Path(
    r"C:\Users\isunderland\everde-ai-operations\scripts\wcro\assets"
    r"\OracleMassUpload_Shell.xlsm"
)
DEFAULT_HD_TEMPLATE = Path(
    r"c:\Users\isunderland\AppData\Local\Temp"
    r"\WK 17 HD FAL_STE SCAL 126 STORES Ship 04-20-26_BM_.xlsm"
)
DEFAULT_SUCCESS_REFERENCE = Path(
    r"c:\Users\isunderland\Downloads\IvanMassUpload.xlsm"
)
DEFAULT_STORE_DRIVEN = Path(
    r"\\192.168.190.10\Claude Sandbox\DataDrops\WCRO"
    r"\_HANDOFF_WCRO_5.47_2026-08-31\reports"
    r"\Store Driven Sales Recommendation"
    r"\HD Store Driven Sales Recommendation (Refresh 5.47 - 2026-08-31).xlsx"
)
DEFAULT_LOW_STORE_DRIVEN = Path(
    r"\\192.168.190.10\Claude Sandbox\DataDrops\WCRO"
    r"\_HANDOFF_WCRO_5.47_2026-08-31\reports"
    r"\Store Driven Sales Recommendation"
    r"\LOW Store Driven Sales Recommendation (Refresh 5.47 - 2026-08-31).xlsx"
)
DEFAULT_OUTPUT = Path(
    r"C:\Users\isunderland\everde-ai-operations\_incoming"
    r"\WCRO_NCA_HD_MassUpload_Test_5.47.xlsm"
)
DEFAULT_XREF_DIR = Path(
    r"\\192.168.190.10\Claude Sandbox\JS Files\Shared\Inventory Cross References"
)

# Address categories that belong to each WCRO market.
# For mass-upload xref gating use CA-core cats only (not AZ/NV/Gulf), so items that
# are only set up for AZ/Vegas are not shipped onto CA store orders.
HD_REGION_ADDR = {
    "N.CA": {"HDNO21", "HDNO29", "HDNO29A"},
    "S.CA": {"HDSO12", "HDSO47", "HDSO48", "HDSO196"},
}
LOW_REGION_ADDR = {
    "N.CA": {"LOWNO"},
    "S.CA": {"LOWSO", "LOWWB"},
}


@dataclass
class ItemCol:
    code: str
    sku: object | None
    desc: str
    src_index: int  # index into WCRO item block (0-based from col H)


def _sku_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def is_citrus_item(code: str, desc: str = "") -> bool:
    """Mass-upload only: citrus stays on WCRO; exclude from landscape Oracle file."""
    code_u = (code or "").strip().upper()
    desc_u = (desc or "").strip().upper()
    if code_u.startswith("CIT"):
        return True
    if "CITRUS" in desc_u:
        return True
    return False


def load_customer_xref_pairs(
    channel: str,
) -> dict[tuple[str, str], set[str]]:
    """
    Return {(sku, item_code): {address_category, ...}} from HD/Lowe's xref.
    Missing file → empty dict (caller may skip filter).
    """
    try:
        from pyxlsb import open_workbook
    except ImportError:
        print("WARN: pyxlsb not installed — skipping customer xref filter")
        return {}

    if channel.upper() == "LOW":
        patterns = ["LOWE'S xref*.xlsb", "LOWE_S_xref*.xlsb", "LOWE*xref*.xlsb"]
    else:
        patterns = [
            "Home Depot Corp-VN=PO xref*.xlsb",
            "Home_Depot_Corp-VN_PO_xref*.xlsb",
            "Home Depot*xref*.xlsb",
        ]

    hits: list[Path] = []
    if DEFAULT_XREF_DIR.is_dir():
        for pat in patterns:
            hits.extend(DEFAULT_XREF_DIR.glob(pat))
    if not hits:
        print(f"WARN: no {channel} xref under {DEFAULT_XREF_DIR}")
        return {}

    path = max(hits, key=lambda p: p.stat().st_mtime)
    pair_addrs: dict[tuple[str, str], set[str]] = {}
    with open_workbook(str(path)) as wb:
        sheet_name = "DATA" if "DATA" in wb.sheets else wb.sheets[0]
        with wb.get_sheet(sheet_name) as sheet:
            rows = [[c.v for c in row] for row in sheet.rows()]
    if not rows:
        return {}
    hdr = [str(h or "").strip() for h in rows[0]]
    try:
        sku_i = hdr.index("SKU")
        item_i = hdr.index("Item")
        addr_i = hdr.index("Address Category")
    except ValueError:
        print(f"WARN: xref missing SKU/Item/Address Category cols: {hdr}")
        return {}

    for r in rows[1:]:
        sku = _sku_str(r[sku_i] if sku_i < len(r) else None)
        item = str((r[item_i] if item_i < len(r) else None) or "").strip()
        addr = str((r[addr_i] if addr_i < len(r) else None) or "").strip()
        if not sku or not item or not addr:
            continue
        # Strip store-suffix variants like CITLIF2101-14795 → keep as-is for exact match;
        # WCRO uses base item codes without -suffix.
        pair_addrs.setdefault((sku, item), set()).add(addr)

    print(f"Loaded {channel} xref {path.name}: {len(pair_addrs)} active SKU×item pairs")
    return pair_addrs


def region_address_categories(channel: str, region: str) -> set[str]:
    table = LOW_REGION_ADDR if channel.upper() == "LOW" else HD_REGION_ADDR
    return set(table.get(region, set()))


def format_store_number(value) -> str:
    """Oracle mass upload expects 4-digit store #s (e.g. 625 → 0625)."""
    if value is None:
        return ""
    raw = str(value).strip()
    if not raw:
        return ""
    # Strip trailing .0 from numeric Excel values
    if raw.endswith(".0") and raw[:-2].isdigit():
        raw = raw[:-2]
    digits = "".join(ch for ch in raw if ch.isdigit())
    if digits:
        return digits.zfill(4)[-4:] if len(digits) <= 4 else digits
    return raw


def format_po_number(po_value, store_4: str):
    """Keep NEEDSPO/{store} in sync with zero-padded store # when present."""
    if po_value is None or str(po_value).strip() == "":
        return f"NEEDSPO/{store_4}" if store_4 else None
    po = str(po_value).strip()
    upper = po.upper()
    if upper.startswith("NEEDSPO/"):
        return f"NEEDSPO/{store_4}" if store_4 else po
    # If PO is just the bare store number, pad it
    digits = "".join(ch for ch in po if ch.isdigit())
    if digits and digits == "".join(ch for ch in str(po_value) if ch.isdigit()):
        if len(digits) <= 4:
            return digits.zfill(4)
    return po_value


def clear_template_used_range(ws, max_row: int = 500, max_col: int = 400) -> None:
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.value is not None:
                cell.value = None


def load_wcro_order_sheet(store_driven_path: Path, region: str) -> list[list]:
    src = openpyxl.load_workbook(str(store_driven_path), data_only=True, read_only=True)
    sheet_name = f"{region} Order"
    if sheet_name not in src.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {src.sheetnames}")
    rows = [list(r) for r in src[sheet_name].iter_rows(values_only=True)]
    src.close()
    return rows


def load_item_master_map(wb) -> dict[str, str]:
    out: dict[str, str] = {}
    if "Item Master" in wb.sheetnames:
        for row in wb["Item Master"].iter_rows(min_row=2, values_only=True):
            if row[0]:
                out[str(row[0]).strip()] = str(row[1]).strip() if row[1] else ""
    elif "SKUs" in wb.sheetnames:
        for row in wb["SKUs"].iter_rows(min_row=2, values_only=True):
            if row[2]:
                out[str(row[2]).strip()] = str(row[3]).strip() if row[3] else ""
    return out


def select_item_columns(
    item_codes,
    skus,
    descs,
    store_rows: list[list],
    item_master_map: dict[str, str],
    channel: str = "HD",
    region: str = "N.CA",
    xref_pairs: dict[tuple[str, str], set[str]] | None = None,
    require_xref: bool = True,
    exclude_citrus: bool = True,
) -> tuple[list[ItemCol], list[str]]:
    """
    Keep only columns that are safe to upload:
      - has Everde item code
      - has a real description (Item Master or WCRO) — never leave #N/A
      - has at least one non-zero store qty
      - unique retailer SKU (when SKU present): if two Everde items share one HD/LOW
        SKU, keep the higher total-qty column and drop the rest
      - NOT citrus (ag restriction — reps order citrus separately)
      - active customer xref for region address categories (SKU×Item pair)
    """
    candidates: list[ItemCol] = []
    dropped: list[str] = []
    n = len(item_codes)
    region_addrs = region_address_categories(channel, region)

    def total_qty(src_j: int) -> float:
        total = 0.0
        for sr in store_rows:
            val = sr[7 + src_j] if 7 + src_j < len(sr) else None
            if val not in (None, 0, ""):
                try:
                    total += float(val)
                except (TypeError, ValueError):
                    pass
        return total

    for j in range(n):
        code = item_codes[j]
        if not code:
            dropped.append(f"col{j}:blank-code")
            continue
        code_s = str(code).strip()
        desc = item_master_map.get(code_s) or (
            str(descs[j]).strip() if j < len(descs) and descs[j] else ""
        )
        if not desc or desc.upper() in ("N/A", "#N/A", "NA"):
            dropped.append(f"{code_s}:no-description")
            continue

        if exclude_citrus and is_citrus_item(code_s, desc):
            dropped.append(f"{code_s}:citrus-ag-restriction")
            continue

        qty = total_qty(j)
        if qty <= 0:
            dropped.append(f"{code_s}:zero-qty")
            continue

        sku_val = skus[j] if j < len(skus) else None
        sku_s = _sku_str(sku_val)

        if require_xref and xref_pairs is not None:
            if not sku_s:
                dropped.append(f"{code_s}:no-sku-for-xref")
                continue
            addrs = xref_pairs.get((sku_s, code_s), set())
            if not addrs:
                dropped.append(f"{code_s}:sku-{sku_s}-not-in-customer-xref")
                continue
            if region_addrs and not (addrs & region_addrs):
                dropped.append(
                    f"{code_s}:sku-{sku_s}-xref-not-in-{region}"
                    f"(addrs={','.join(sorted(addrs)[:4])})"
                )
                continue

        candidates.append(
            ItemCol(
                code=code_s,
                sku=sku_val,
                desc=desc,
                src_index=j,
            )
        )

    # Deduplicate by retailer SKU — keep highest total qty
    by_sku: dict[str, list[ItemCol]] = {}
    no_sku: list[ItemCol] = []
    for item in candidates:
        if item.sku is None or str(item.sku).strip() == "":
            no_sku.append(item)
            continue
        key = _sku_str(item.sku)
        by_sku.setdefault(key, []).append(item)

    kept: list[ItemCol] = list(no_sku)
    for sku_key, group in by_sku.items():
        if len(group) == 1:
            kept.append(group[0])
            continue
        ranked = sorted(group, key=lambda it: total_qty(it.src_index), reverse=True)
        kept.append(ranked[0])
        for loser in ranked[1:]:
            dropped.append(
                f"{loser.code}:duplicate-sku-{sku_key}-kept-{ranked[0].code}"
            )

    # Preserve original WCRO column order
    kept.sort(key=lambda it: it.src_index)
    return kept, dropped


def validate_locked_format(ws, n_items: int, n_stores: int) -> list[str]:
    """Return list of contract violations (empty = OK)."""
    errors: list[str] = []
    last_row = 13 + n_stores

    # Row 7 must be empty
    for col in range(1, 8 + n_items + 1):
        if ws.cell(7, col).value is not None:
            errors.append(f"Row7 not blank at col {col}: {ws.cell(7, col).value!r}")
            break

    # E/F blank on header + data
    for r in range(13, last_row + 1):
        if ws.cell(r, 5).value is not None or ws.cell(r, 6).value is not None:
            errors.append(f"E/F must be blank at row {r}")
            break

    # No formulas / #N/A on row 12 item band
    for j in range(n_items):
        v = ws.cell(12, 8 + j).value
        if v is None:
            errors.append(f"Row12 missing description at col {get_column_letter(8+j)}")
            continue
        s = str(v)
        if s.startswith("=") or "VLOOKUP" in s.upper() or s.strip().upper() in (
            "N/A",
            "#N/A",
        ):
            errors.append(f"Row12 bad value at {get_column_letter(8+j)}: {s[:60]!r}")

    # No long substituting notes on row 10
    for j in range(n_items):
        v = ws.cell(10, 8 + j).value
        if v and "substitut" in str(v).lower():
            errors.append(f"Row10 still has substituting note at {get_column_letter(8+j)}")

    # Store # must be 4-digit zero-padded text
    for r in range(14, last_row + 1):
        s = ws.cell(r, 2).value
        if s is None:
            continue
        ss = str(s).strip()
        if not (len(ss) >= 4 and ss.isdigit()):
            errors.append(f"Store # at row {r} must be 4-digit text, got {ss!r}")
            break

    # Required labels
    if ws.cell(1, 1).value != "Customer Name":
        errors.append("A1 label missing")
    if ws.cell(13, 1).value != "Rep" or ws.cell(13, 7).value != "Internal # ":
        errors.append("Row13 labels incorrect")

    return errors


def build(
    store_driven_path: Path,
    oracle_template_path: Path,
    output_path: Path,
    region: str = "N.CA",
    channel: str = "HD",
    customer_name: str = "HOME DEPOT CORP - VN",
    address_category: str = "",
    order_type: str = "",
    req_delivery_date=None,
    ship_instr: str | None = None,
    require_xref: bool = True,
    exclude_citrus: bool = True,
) -> None:
    if not oracle_template_path.exists():
        raise FileNotFoundError(f"Oracle template not found: {oracle_template_path}")
    if not store_driven_path.exists():
        raise FileNotFoundError(f"WCRO Store Driven not found: {store_driven_path}")

    src_rows = load_wcro_order_sheet(store_driven_path, region)

    header_row = src_rows[12]
    sku_row = src_rows[10]
    desc_row = src_rows[11]

    raw_codes = list(header_row[7:])
    raw_skus = list(sku_row[7:])
    raw_descs = list(desc_row[7:])
    n_raw = len([c for c in raw_codes if c is not None])
    raw_codes = raw_codes[:n_raw]
    raw_skus = (raw_skus + [None] * n_raw)[:n_raw]
    raw_descs = (raw_descs + [None] * n_raw)[:n_raw]

    store_rows = [r for r in src_rows[13:] if r[1] is not None]
    if not store_rows:
        raise ValueError(f"No store rows found on '{region} Order'")

    # Stage shell early so Item Master is available for column selection
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()
    shutil.copy2(oracle_template_path, output_path)

    wb = openpyxl.load_workbook(str(output_path), keep_vba=True)
    if "Template" not in wb.sheetnames:
        raise ValueError(f"Expected 'Template' sheet in {oracle_template_path.name}")
    if not wb.vba_archive:
        raise RuntimeError("VBA archive missing — Upload To Oracle would not work.")

    item_master_map = load_item_master_map(wb)
    xref_pairs = load_customer_xref_pairs(channel) if require_xref else {}
    if require_xref and not xref_pairs:
        print(
            "WARN: xref filter requested but no pairs loaded — "
            "continuing without xref gate this run"
        )
        require_xref_eff = False
    else:
        require_xref_eff = require_xref

    items, dropped = select_item_columns(
        raw_codes,
        raw_skus,
        raw_descs,
        store_rows,
        item_master_map,
        channel=channel,
        region=region,
        xref_pairs=xref_pairs if require_xref_eff else None,
        require_xref=require_xref_eff,
        exclude_citrus=exclude_citrus,
    )
    if not items:
        raise ValueError("No upload-safe item columns remained after filtering")

    n_items = len(items)
    print(f"Oracle shell: {oracle_template_path.name}")
    print(f"Channel/region: {channel} {region}")
    print(f"Stores: {len(store_rows)}")
    print(f"Items kept: {n_items} (dropped {len(dropped)} unsafe/empty columns)")
    citrus_drops = [d for d in dropped if "citrus" in d]
    xref_drops = [d for d in dropped if "xref" in d]
    if citrus_drops:
        print(f"  citrus excluded: {len(citrus_drops)}")
    if xref_drops:
        print(f"  xref excluded: {len(xref_drops)}")
    if dropped:
        for d in dropped[:12]:
            print(f"  drop: {d}")
        if len(dropped) > 12:
            print(f"  ... +{len(dropped) - 12} more")

    # Drop stores that have no remaining qty after item filters
    kept_store_rows: list[list] = []
    for sr in store_rows:
        has_qty = False
        for item in items:
            val = sr[7 + item.src_index] if 7 + item.src_index < len(sr) else None
            if val not in (None, 0, ""):
                has_qty = True
                break
        if has_qty:
            kept_store_rows.append(sr)
    if len(kept_store_rows) < len(store_rows):
        print(
            f"Stores with qty after filters: {len(kept_store_rows)} "
            f"(dropped {len(store_rows) - len(kept_store_rows)} empty)"
        )
    store_rows = kept_store_rows
    if not store_rows:
        raise ValueError("No stores left with quantities after item filters")

    n_items = len(items)
    ws = wb["Template"]
    clear_template_used_range(ws)

    hdr_font = Font(bold=True, size=11)
    title_font = Font(bold=True, size=12)
    thin = Border(
        left=Side("thin"),
        right=Side("thin"),
        top=Side("thin"),
        bottom=Side("thin"),
    )
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    item_fill = PatternFill("solid", fgColor="FCE4D6")

    # --- Header (B2/B4 left for user or Teams bot) ---
    ws.cell(1, 1, "Customer Name").font = hdr_font
    ws.cell(1, 2, customer_name).font = title_font

    ws.cell(2, 1, "Req. Delivery Date").font = hdr_font
    if req_delivery_date:
        ws.cell(2, 2, req_delivery_date)

    ws.cell(3, 1, "Shipping Instr.").font = hdr_font
    ws.cell(
        3,
        2,
        ship_instr
        or f"{region.replace('.', '')} SPREAD - {len(store_rows)} STORES",
    )

    ws.cell(4, 1, "Order Type").font = hdr_font
    if order_type:
        ws.cell(4, 2, order_type)
    ws.cell(4, 5, " ")  # spacer used by some shells

    ws.cell(5, 1, "# of materials").font = hdr_font
    ws.cell(5, 2, n_items)

    ws.cell(6, 1, "Address Category").font = hdr_font
    if address_category:
        ws.cell(6, 2, address_category)

    # Row 7 intentionally blank

    last_data_row = 14 + len(store_rows) - 1

    for j in range(n_items):
        col = 8 + j
        cl = get_column_letter(col)
        cell = ws.cell(8, col, f"=SUM({cl}14:{cl}{last_data_row})")
        cell.font = Font(bold=True, size=10)
        cell.fill = green_fill

    # Row 10 — label only (no substituting text in the upload band)
    ws.cell(10, 7, "Picking Notes").font = hdr_font

    # Row 11 — SKUs
    ws.cell(11, 7, "Sku # ").font = hdr_font
    for j, item in enumerate(items):
        if item.sku is not None and str(item.sku).strip() != "":
            cell = ws.cell(11, 8 + j, item.sku)
            cell.font = Font(size=9)
            cell.fill = item_fill

    # Row 12 — static descriptions only
    ws.cell(12, 7, "Item").font = hdr_font
    for j, item in enumerate(items):
        ws.cell(12, 8 + j, item.desc).font = Font(size=8)

    # Row 13 — E/F blank by contract
    for i, h in enumerate(["Rep", "Store #", "City", "PO#"]):
        cell = ws.cell(13, i + 1, h)
        cell.font = hdr_font
        cell.fill = header_fill
        cell.border = thin
    cell = ws.cell(13, 7, "Internal # ")
    cell.font = hdr_font
    cell.fill = header_fill
    cell.border = thin

    for j, item in enumerate(items):
        cell = ws.cell(13, 8 + j, item.code)
        cell.font = Font(bold=True, size=8)
        cell.fill = header_fill
        cell.border = thin

    # Store grid — E/F blank; Store # always 4-digit zero-padded text
    for si, sr in enumerate(store_rows):
        rn = 14 + si
        store_4 = format_store_number(sr[1])
        ws.cell(rn, 1, sr[0])  # Rep
        store_cell = ws.cell(rn, 2, store_4)
        store_cell.number_format = "@"  # force text so leading zero is kept
        ws.cell(rn, 3, sr[2])  # City
        ws.cell(rn, 4, format_po_number(sr[3], store_4))  # PO#
        for j, item in enumerate(items):
            src_j = item.src_index
            val = sr[7 + src_j] if 7 + src_j < len(sr) else None
            if val not in (None, 0, ""):
                cell = ws.cell(rn, 8 + j, val)
                cell.border = thin
        for ci in (1, 2, 3, 4, 7):
            ws.cell(rn, ci).border = thin

    widths = {"A": 22, "B": 10, "C": 20, "D": 18, "E": 12, "F": 12, "G": 14}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w
    for j in range(n_items):
        ws.column_dimensions[get_column_letter(8 + j)].width = 12

    violations = validate_locked_format(ws, n_items, len(store_rows))
    if violations:
        wb.close()
        raise RuntimeError(
            "Locked format validation failed:\n  - " + "\n  - ".join(violations)
        )

    wb.save(str(output_path))
    wb.close()

    print(f"Saved: {output_path}")
    print("Format OK: blank E/F, blank row7, no #N/A, B2/B4 ready for user/bot")
    if not req_delivery_date or not order_type:
        print("Fill before upload: B2 Req. Delivery Date, B4 Order Type")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--store-driven", default="",
                    help="WCRO Store Driven workbook (default from --channel)")
    ap.add_argument(
        "--template",
        default=str(DEFAULT_ORACLE_TEMPLATE),
        help="Oracle .xlsm shell (default: IvanMassUpload-derived golden shell)",
    )
    ap.add_argument(
        "--prefer-hd-shell",
        action="store_true",
        help="Use Temp WK17 HD SCAL shell instead of golden IvanMassUpload shell",
    )
    ap.add_argument("--channel", default="HD", choices=["HD", "LOW"])
    ap.add_argument("--region", default="N.CA")
    ap.add_argument("--output", default=str(DEFAULT_OUTPUT))
    ap.add_argument("--customer", default="")
    ap.add_argument("--address-category", default="")
    ap.add_argument("--order-type", default="")
    ap.add_argument("--req-delivery-date", default="")
    ap.add_argument("--ship-instr", default="")
    ap.add_argument(
        "--skip-xref-filter",
        action="store_true",
        help="Do not require active customer xref for region address categories",
    )
    ap.add_argument(
        "--include-citrus",
        action="store_true",
        help="Keep citrus items (default: exclude — ag restriction, order separately)",
    )
    args = ap.parse_args()

    template = Path(args.template)
    if args.prefer_hd_shell and DEFAULT_HD_TEMPLATE.exists():
        template = DEFAULT_HD_TEMPLATE
    elif not template.exists() and DEFAULT_SUCCESS_REFERENCE.exists():
        template = DEFAULT_SUCCESS_REFERENCE
    elif not template.exists() and DEFAULT_HD_TEMPLATE.exists():
        template = DEFAULT_HD_TEMPLATE

    store_driven = Path(args.store_driven) if args.store_driven.strip() else (
        DEFAULT_LOW_STORE_DRIVEN if args.channel == "LOW" else DEFAULT_STORE_DRIVEN
    )
    customer = args.customer.strip() or (
        "LOWE''S COMPANIES, INC. - VN"
        if args.channel == "LOW"
        else "HOME DEPOT CORP - VN"
    )

    req_date = args.req_delivery_date.strip() or None
    if req_date:
        req_date = datetime.strptime(req_date, "%Y-%m-%d")

    build(
        store_driven,
        template,
        Path(args.output),
        region=args.region,
        channel=args.channel,
        customer_name=customer,
        address_category=args.address_category,
        order_type=args.order_type.strip(),
        req_delivery_date=req_date,
        ship_instr=args.ship_instr.strip() or None,
        require_xref=not args.skip_xref_filter,
        exclude_citrus=not args.include_citrus,
    )
